#!/usr/bin/env python3
"""Generate Anthony Ty Marler's professional resume as DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BODY = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
RULE = "1F3A5F"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=BODY):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_run_font_ascii(run, name="Calibri"):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def add_bottom_border(paragraph, color=RULE, size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is not None:
        pPr.remove(pBdr)
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_tab_stop(paragraph, position_inches, alignment=WD_TAB_ALIGNMENT.RIGHT):
    pPr = paragraph._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        pPr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right" if alignment == WD_TAB_ALIGNMENT.RIGHT else "left")
    tab.set(qn("w:pos"), str(int(position_inches * 1440)))
    tabs.append(tab)


def heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text.upper())
    set_run_font(run, size=11, bold=True, color=NAVY)
    set_run_font_ascii(run)
    add_bottom_border(p)
    return p


def body_para(doc, text, size=10, italic=False, space_after=3, space_before=0, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=size, italic=italic, bold=bold, color=BODY)
    set_run_font_ascii(run)
    return p


def job_header(doc, company, location, title, dates, right_margin_in=7.2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    add_tab_stop(p, right_margin_in)
    run = p.add_run(company)
    set_run_font(run, size=11, bold=True, color=NAVY)
    set_run_font_ascii(run)
    if location:
        run = p.add_run(f"  |  {location}")
        set_run_font(run, size=10, color=MUTED)
        set_run_font_ascii(run)
    run = p.add_run("\t" + dates)
    set_run_font(run, size=10, italic=True, color=MUTED)
    set_run_font_ascii(run)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.05
    run = p2.add_run(title)
    set_run_font(run, size=10.5, italic=True, bold=True, color=BODY)
    set_run_font_ascii(run)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.space_before = Pt(0.5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.left_indent = Inches(0.25 + 0.15 * level)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    run = p.add_run(text)
    set_run_font(run, size=10, color=BODY)
    set_run_font_ascii(run)
    return p


def compact_role(doc, company, location, title, dates, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    add_tab_stop(p, 7.2)
    run = p.add_run(company)
    set_run_font(run, size=10.5, bold=True, color=NAVY)
    set_run_font_ascii(run)
    if location:
        run = p.add_run(f"  |  {location}")
        set_run_font(run, size=10, color=MUTED)
        set_run_font_ascii(run)
    run = p.add_run("\t" + dates)
    set_run_font(run, size=10, italic=True, color=MUTED)
    set_run_font_ascii(run)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(1)
    p2.paragraph_format.line_spacing = 1.08
    run = p2.add_run(title + ".  ")
    set_run_font(run, size=10, italic=True, bold=True, color=BODY)
    set_run_font_ascii(run)
    run = p2.add_run(text)
    set_run_font(run, size=10, color=BODY)
    set_run_font_ascii(run)


def skill_line(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(label + ":  ")
    set_run_font(run, size=10, bold=True, color=NAVY)
    set_run_font_ascii(run)
    run = p.add_run(text)
    set_run_font(run, size=10, color=BODY)
    set_run_font_ascii(run)


def disable_spellcheck(doc):
    settings = doc.settings.element
    hide = OxmlElement("w:hideSpellingErrors")
    hide.set(qn("w:val"), "true")
    settings.append(hide)
    hide2 = OxmlElement("w:hideGrammaticalErrors")
    hide2.set(qn("w:val"), "true")
    settings.append(hide2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.45)
    disable_spellcheck(doc)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")

    # Tighten built-in list bullet
    try:
        bullet = styles["List Bullet"]
        bullet.font.name = "Calibri"
        bullet.font.size = Pt(10.5)
        pf = bullet.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(1)
    except KeyError:
        pass

    # Header
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(0)
    name.paragraph_format.space_after = Pt(2)
    name.paragraph_format.line_spacing = 1.0
    run = name.add_run("ANTHONY TY MARLER")
    set_run_font(run, size=20, bold=True, color=NAVY)
    set_run_font_ascii(run)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_before = Pt(0)
    contact.paragraph_format.space_after = Pt(2)
    contact.paragraph_format.line_spacing = 1.05
    run = contact.add_run(
        "Winnsboro, Texas 75494  ·  (903) 474-3647  ·  tymarler@gmail.com"
    )
    set_run_font(run, size=10.5, color=BODY)
    set_run_font_ascii(run)

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag.paragraph_format.space_before = Pt(0)
    tag.paragraph_format.space_after = Pt(2)
    tag.paragraph_format.line_spacing = 1.0
    run = tag.add_run(
        "Senior Application Architect  ·  Inventor  ·  Energy, Climate-Tech & AI-Assisted Systems"
    )
    set_run_font(run, size=10.5, italic=True, color=NAVY)
    set_run_font_ascii(run)
    add_bottom_border(tag, size="18")

    heading(doc, "Professional Summary")
    body_para(
        doc,
        "Results-oriented technical strategist and application architect who builds mission-critical "
        "software and data-driven systems for operational efficiency, regulatory compliance, and "
        "business outcomes. Named inventor on two issued U.S. patents. Combines AWS cloud, .NET, "
        "SQL, and systems integration with generative AI (Cursor, Claude, ChatGPT) and deep expertise "
        "in energy, mining, manufacturing, and environmental systems. Leads cross-functional teams "
        "through the full project lifecycle.",
        space_after=2,
    )

    heading(doc, "Selected Highlights")
    add_bullet(
        doc,
        "AI-built applications — Used Cursor, Claude, and ChatGPT to architect and develop an "
        "international fugitive-tracking application, a custom architectural-planning site, a "
        "clothing-brand retail site, and an internet television streaming application.",
    )
    add_bullet(
        doc,
        "Vistra — Designed a data-driven process engine for dynamic tasks; mine planning and "
        "monitoring software that measures productivity by shift; and a cloud emissions platform "
        "that unified diverse systems for credit trading, accounting for $60 million in profit "
        "in the first two months.",
    )
    add_bullet(
        doc,
        "Compaq — Invented and developed a patented software-delivery system that eliminated more "
        "than 70% of computer manufacturing time, plus a post-sales configuration method that "
        "increased inventory flexibility.",
    )
    add_bullet(
        doc,
        "Texas Built Biofuels — Invented a process-intensified biodiesel method that reduced "
        "reaction time from 4 hours to 4 minutes (60x throughput; ~40% lower energy cost per "
        "gallon; ASTM-quality fuel).",
    )
    add_bullet(
        doc,
        "Automated Response Information Systems — Designed and implemented a hardware-based "
        "voice-recognition system for hands-free data collection in harsh environments "
        "(early applied AI / speech recognition).",
    )

    heading(doc, "Core Competencies")
    body_para(
        doc,
        "Application Architecture  ·  Generative AI (Cursor, Claude, ChatGPT)  ·  Systems Integration  ·  "
        "AWS Cloud & Chef  ·  .NET / C# / SQL  ·  Agile / Scrum  ·  Cross-Functional Leadership  ·  "
        "Root-Cause Analysis  ·  TCO Analysis  ·  Data Modeling & Database Design  ·  Emissions "
        "Compliance & Carbon Credit Trading  ·  Mining Operations Software  ·  GIS  ·  SharePoint  ·  "
        "OSIsoft PI  ·  Tableau",
        space_after=2,
    )

    heading(doc, "Professional Experience")

    job_header(
        doc,
        "Independent",
        "Texas",
        "Application Architect, AI-Assisted Systems",
        "Mar 2025 – Present",
    )
    add_bullet(
        doc,
        "Used Cursor, Claude, and ChatGPT to architect and develop an international "
        "fugitive-tracking application and a custom architectural-planning website.",
    )
    add_bullet(
        doc,
        "Delivered a clothing-brand retail / e-commerce site and an internet television "
        "streaming application with AI-assisted architecture, UI, and backend development.",
    )

    job_header(
        doc,
        "Vistra Corporate Services Company",
        "Texas",
        "Senior Analyst / Architect, Generation",
        "Dec 2019 – Feb 2025",
    )
    add_bullet(
        doc,
        "Designed and developed a data-driven process engine to manage dynamic tasks, providing "
        "automated steps to drive the process to completion.",
    )
    add_bullet(
        doc,
        "Designed and developed mine planning and monitoring software to calculate productivity "
        "and measure results by shift.",
    )
    add_bullet(
        doc,
        "Designed and developed a cloud-based AWS platform that ingested SCADA, PLC, and "
        "third-party emissions data into a normalized database for credit trading; accounted "
        "for $60 million in profit in the first two months, cut trading latency from weeks to "
        "hours, and reduced compliance reporting time by 80%.",
    )
    add_bullet(
        doc,
        "Led web applications supporting mining, fossil generation, battery, and solar operations, "
        "including Tableau dashboards for traders and compliance officers.",
    )

    job_header(
        doc,
        "Accenture",
        "Texas",
        "Business Application Manager",
        "Aug 2016 – Dec 2019",
    )
    add_bullet(
        doc,
        "Led day-to-day maintenance, enhancement, and operations for the Fossil Applications "
        "portfolio serving Luminant generation and mining.",
    )
    add_bullet(
        doc,
        "Directed development across AWS, Chef, .NET, C#, C++, PHP, SQL Server, MySQL, Tableau, "
        "Formotus, TDMobile, StackVision, and NetDAHS.",
    )
    add_bullet(
        doc,
        "Owned troubleshooting and RCA for production issues across ROTT, Skaief, SmartProcedures, "
        "GIS, Mobility, Safety Index, Proact, Digital Signage, Qdabra, SharePoint, PI, and land management.",
    )
    add_bullet(
        doc,
        "Led a cloud modernization that aligned DevOps, QA, and engineering and delivered three "
        "weeks ahead of schedule. Mentored engineers on AWS and Chef; a junior developer became "
        "a lead contributor within six months.",
    )

    job_header(
        doc,
        "Capgemini, LLC",
        "Texas",
        "Senior Application Architect",
        "Feb 2008 – Aug 2016",
    )
    add_bullet(
        doc,
        "Architected and delivered Fossil Generation applications covering MSHA training, "
        "emissions compliance, procedures, GIS, safety, mobility, digital signage, and SharePoint.",
    )
    add_bullet(
        doc,
        "Provided architectural solutions for new builds and rebuilds to improve reliability "
        "and functionality, including TCO analysis of existing and future-state environments.",
    )
    add_bullet(
        doc,
        "Served as SME for mining operations, mining software systems, data modeling, and "
        "database design; introduced mobility (Formotus, TDMobile), OSIsoft PI, land management, "
        "and application monitoring.",
    )

    heading(doc, "Additional Experience")
    compact_role(
        doc,
        "Texas Built Company, LLC",
        "Winnsboro, TX",
        "Owner",
        "Feb 1998 – Feb 2008",
        "Invented a process-intensified biodiesel reactor that reduced reaction time from four "
        "hours to four minutes, enabling continuous rather than batch production, about 60x "
        "throughput in the same footprint, and roughly 40% lower energy cost per gallon while "
        "maintaining ASTM fuel quality. Patent pending: Texas Built process intensified biofuels "
        "reactor technologies.",
    )
    compact_role(
        doc,
        "Samsung / AST Research",
        "",
        "Principal Engineer / Worldwide Software Systems Architect",
        "Mar 1997 – Feb 1998",
        "Worldwide software systems architecture for PC manufacturing and configuration platforms.",
    )
    compact_role(
        doc,
        "Compaq Computer Corporation",
        "Houston, TX",
        "Senior Systems Engineer",
        "Oct 1993 – Mar 1997",
        "Invented and developed a patented software-delivery system that eliminated more than "
        "70% of the manufacture time of their computers. Also invented a post-sales configuration "
        "method that increased inventory flexibility. Named inventor on U.S. Patents 6,202,070 "
        "and 5,974,567.",
    )
    compact_role(
        doc,
        "Automated Response Information Systems, Inc.",
        "",
        "Owner",
        "Oct 1991 – Oct 1993",
        "Designed and implemented a hardware-based voice-recognition system for hands-free "
        "data collection in harsh environments (early applied AI / speech recognition).",
    )
    compact_role(
        doc,
        "Stanford Telecommunications, Inc.",
        "",
        "Systems Engineer",
        "Jul 1989 – Aug 1991",
        "Worked on a team to develop planning and control software for the defense communication "
        "satellite network.",
    )
    compact_role(
        doc,
        "Elcom and Associates",
        "",
        "Field Engineer",
        "Feb 1987 – Jul 1989",
        "Field engineering and technical support for communications and electronics systems.",
    )
    compact_role(
        doc,
        "U.S. Army",
        "",
        "Strategic Microwave System Repair Technician",
        "Feb 1984 – Feb 1987",
        "Maintained strategic microwave communications; completed Army electronics, microwave "
        "repair, and cryptographic training.",
    )

    heading(doc, "Patents")
    add_bullet(
        doc,
        "U.S. 6,202,070 (2001) — Computer manufacturing system architecture with enhanced "
        "software distribution functions (Compaq; named inventor).",
    )
    add_bullet(
        doc,
        "U.S. 5,974,567 (1999) — Ghost partition / runtime-selectable file system "
        "(Compaq; named inventor).",
    )
    add_bullet(
        doc,
        "Pending — Texas Built process intensified biofuels reactor technologies.",
    )

    heading(doc, "Technical Skills")
    skill_line(
        doc,
        "AI & developer tools",
        "Cursor, Claude, ChatGPT; AI-assisted architecture, coding, RCA, and documentation",
    )
    skill_line(
        doc,
        "Languages & platforms",
        "C# / .NET, VB.NET, C++ / C++.NET, C, PHP, Java, SQL, AJAX, Chef, Intel Assembly; "
        "prior: Ada, FORTRAN, BASIC, VMS/DCL, PRO*C, FOCUS, PAL, Silverlight",
    )
    skill_line(
        doc,
        "Cloud, data & monitoring",
        "AWS, Amazon RDS, Chef, APIs, OSIsoft PI, MS SQL Server, MySQL, Oracle, DB2, Tableau, "
        "SCADA/PLC integration, BAM application monitoring, StackVision, NetDAHS",
    )
    skill_line(
        doc,
        "Applications & tools",
        "Visual Studio, SharePoint, ESRI GIS, Formotus, TDMobile, Qdabra, SmartProcedures, "
        "Visio, MS Project, LabVIEW",
    )
    skill_line(
        doc,
        "Domains",
        "Fossil generation, mining operations, emissions compliance and carbon credit trading, "
        "MSHA training, e-commerce, streaming media, architectural planning systems, "
        "investigative/tracking applications, biofuels process intensification, "
        "AI-assisted industrial systems",
    )

    heading(doc, "Education & Professional Development")
    body_para(
        doc,
        "Computer Science coursework and Certificate in Ada, Colorado Technical College. "
        "Advanced C++ and Advanced Java, Compaq. DEC VAX/VMS System Management. U.S. Army "
        "Basic Electronics, Strategic Microwave System Repair, and Cryptographic courses.",
        space_after=0,
    )

    out = Path(__file__).resolve().parent / "Anthony_Ty_Marler_Resume.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build()
