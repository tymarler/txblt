#!/usr/bin/env python3
"""Generate an Indeed / ATS upload resume for Anthony Ty Marler.

Source of truth is the latest combined resume on cursor/resume-generation-28e1
(commit 7233acf, reverse-chronological) plus later LinkedIn copy:
Fractional CTO in the headline, Army Feb 1984-Feb 1987 with the applied
course names, and the four Independent AI-built applications.

Indeed / ATS rules applied here:
- .docx is the primary upload format
- single column; no tables, text boxes, headers, footers, or images
- dates on their own line (no right-aligned tabs)
- standard headings: Summary, Skills, Experience, Education, Patents
- full month names and ASCII hyphens
"""

from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BODY = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)

NAME = "Anthony Ty Marler"
HEADLINE = (
    "Senior Application Architect | Inventor | Energy and Climate-Tech | "
    "AI-Assisted Systems | Fractional CTO"
)
LOCATION = "Winnsboro, TX 75494"
PHONE = "(903) 474-3647"
EMAIL = "tymarler@gmail.com"

SUMMARY = (
    "Senior application architect and fractional CTO currently shipping "
    "AI-assisted products with Cursor, Claude, and ChatGPT. Previously spent "
    "five years at Vistra designing generation, mining, and emissions systems, "
    "including a cloud credit-trading platform that accounted for $60 million "
    "in profit in the first two months. Named inventor on two U.S. patents. "
    "Combines AWS, .NET, SQL, and systems integration with energy, mining, and "
    "manufacturing domain expertise, and leads cross-functional teams through "
    "the full project lifecycle."
)

# One keyword block parses more reliably on Indeed than labeled subgroups.
SKILLS = (
    "Application Architecture, Generative AI, Cursor, Claude, ChatGPT, "
    "Fractional CTO, Systems Integration, Amazon Web Services (AWS), Chef, "
    "Microsoft .NET, C#, VB.NET, C++, SQL, PHP, Java, Agile, Scrum, Cloud "
    "Computing, Data Modeling, Database Design, Root-Cause Analysis, TCO "
    "Analysis, Cross-Functional Leadership, Tableau, SharePoint, OSIsoft PI, "
    "GIS, ESRI GIS, SCADA, PLC, Amazon RDS, MS SQL Server, MySQL, Oracle, "
    "DB2, APIs, StackVision, NetDAHS, Formotus, TDMobile, Visual Studio, "
    "Emissions Compliance, Carbon Credit Trading, Mining Operations Software, "
    "MSHA Training, Biofuels Process Intensification"
)

# Newest first. Dates use full month names for Indeed parsing.
JOBS = [
    {
        "title": "Application Architect, AI-Assisted Systems",
        "company": "Independent",
        "location": "Texas",
        "dates": "March 2025 - Present",
        "bullets": [
            "Used Cursor, Claude, and ChatGPT to architect and develop an international fugitive-tracking application and a custom architectural-planning website.",
            "Delivered a clothing-brand retail / e-commerce site and an internet television streaming application with AI-assisted architecture, UI, and backend development.",
        ],
    },
    {
        "title": "Senior Analyst / Architect, Generation",
        "company": "Vistra Corporate Services Company",
        "location": "Texas",
        "dates": "December 2019 - February 2025",
        "bullets": [
            "Designed and developed a data-driven process engine to manage dynamic tasks, providing automated steps to drive the process to completion.",
            "Designed and developed mine planning and monitoring software to calculate productivity and measure results by shift.",
            "Designed and developed a cloud-based AWS platform that ingested SCADA, PLC, and third-party emissions data into a normalized database for credit trading; accounted for $60 million in profit in the first two months, cut trading latency from weeks to hours, and reduced compliance reporting time by 80%.",
            "Led web applications supporting mining, fossil generation, battery, and solar operations, including Tableau dashboards for traders and compliance officers.",
        ],
    },
    {
        "title": "Business Application Manager",
        "company": "Accenture",
        "location": "Texas",
        "dates": "August 2016 - December 2019",
        "bullets": [
            "Led day-to-day maintenance, enhancement, and operations for the Fossil Applications portfolio serving Luminant generation and mining.",
            "Directed development across AWS, Chef, .NET, C#, C++, PHP, SQL Server, MySQL, Tableau, Formotus, TDMobile, StackVision, and NetDAHS.",
            "Owned troubleshooting and root-cause analysis for production issues across ROTT, Skaief, SmartProcedures, GIS, Mobility, Safety Index, Proact, Digital Signage, Qdabra, SharePoint, PI, and land management.",
            "Led the Fossil Applications cloud modernization on AWS and Chef, putting DevOps, QA, and engineering on one release path and finishing the cutover ahead of schedule.",
        ],
    },
    {
        "title": "Senior Application Architect",
        "company": "Capgemini, LLC",
        "location": "Texas",
        "dates": "February 2008 - August 2016",
        "bullets": [
            "Architected and delivered Fossil Generation applications covering MSHA training, emissions compliance, procedures, GIS, safety, mobility, digital signage, and SharePoint.",
            "Provided architectural solutions for new builds and rebuilds to improve reliability and functionality, including TCO analysis of existing and future-state environments.",
            "Served as SME for mining operations, mining software systems, data modeling, and database design; introduced mobility (Formotus, TDMobile), OSIsoft PI, land management, and application monitoring.",
        ],
    },
    {
        "title": "Owner",
        "company": "Texas Built Company, LLC",
        "location": "Winnsboro, TX",
        "dates": "February 1998 - February 2008",
        "bullets": [
            "Invented a process-intensified biodiesel reactor that reduced reaction time from four hours to four minutes, enabling continuous rather than batch production, about 60x throughput in the same footprint, and roughly 40% lower energy cost per gallon while maintaining ASTM fuel quality.",
            "Patent pending: Texas Built process intensified biofuels reactor technologies.",
        ],
    },
    {
        "title": "Principal Engineer / Worldwide Software Systems Architect",
        "company": "Samsung / AST Research",
        "location": "United States",
        "dates": "March 1997 - February 1998",
        "bullets": [
            "Worldwide software systems architecture for PC manufacturing and configuration platforms.",
        ],
    },
    {
        "title": "Senior Systems Engineer",
        "company": "Compaq Computer Corporation",
        "location": "Houston, TX",
        "dates": "October 1993 - March 1997",
        "bullets": [
            "Invented and developed a patented software-delivery system that eliminated more than 70% of the manufacture time of Compaq computers. Also invented a post-sales configuration method that increased inventory flexibility.",
            "Named inventor on U.S. Patents 6,202,070 and 5,974,567.",
        ],
    },
    {
        "title": "Owner",
        "company": "Automated Response Information Systems, Inc.",
        "location": "United States",
        "dates": "October 1991 - October 1993",
        "bullets": [
            "Designed and implemented a hardware-based voice-recognition system for hands-free data collection in harsh environments (early applied AI / speech recognition).",
        ],
    },
    {
        "title": "Systems Engineer",
        "company": "Stanford Telecommunications, Inc.",
        "location": "United States",
        "dates": "July 1989 - August 1991",
        "bullets": [
            "Worked on a team to develop planning and control software for the defense communication satellite network.",
        ],
    },
    {
        "title": "Field Engineer",
        "company": "Elcom and Associates",
        "location": "United States",
        "dates": "February 1987 - July 1989",
        "bullets": [
            "Field engineering and technical support for communications and electronics systems.",
        ],
    },
    {
        "title": "Strategic Microwave System Repair Technician",
        "company": "U.S. Army",
        "location": "United States",
        "dates": "February 1984 - February 1987",
        "bullets": [
            "Maintained strategic microwave communications systems. Completed U.S. Army Basic Electronics, Communications Electronics Strategic Microwave System Repair, and Cryptographic courses.",
        ],
    },
]

PATENTS = [
    "Pending - Texas Built process intensified biofuels reactor technologies.",
    "U.S. 6,202,070 (issued March 13, 2001) - Computer manufacturing system architecture with enhanced software distribution functions (Compaq Computer Corporation; named inventor).",
    "U.S. 5,974,567 (issued October 26, 1999) - Ghost partition / runtime-selectable file system (Compaq Computer Corporation; named inventor).",
]

EDUCATION = [
    {
        "school": "Colorado Technical College",
        "detail": "Computer Science coursework and Certificate in Ada",
    },
    {
        "school": "Compaq Computer Corporation",
        "detail": "Advanced C++ and Advanced Java",
    },
    {
        "school": "Professional Development",
        "detail": "DEC VAX/VMS System Management",
    },
    {
        "school": "U.S. Army",
        "detail": "Basic Electronics; Communications Electronics Strategic Microwave System Repair; Cryptographic courses",
    },
]


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=BODY):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def add_bottom_border(paragraph, color="1F3A5F", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def disable_spellcheck(doc):
    settings = doc.settings.element
    hide = OxmlElement("w:hideSpellingErrors")
    hide.set(qn("w:val"), "true")
    settings.append(hide)
    hide2 = OxmlElement("w:hideGrammaticalErrors")
    hide2.set(qn("w:val"), "true")
    settings.append(hide2)


def add_para(doc, text, *, size=10, bold=False, italic=False, color=BODY, align="left",
             before=0, after=2, space=1.08):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = space
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text):
    p = add_para(doc, text.upper(), size=11, bold=True, color=NAVY, before=7, after=2, space=1.0)
    add_bottom_border(p)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    run = p.add_run(text)
    set_run_font(run, size=10, color=BODY)
    return p


def build_docx() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.45)
    # No header/footer content - Indeed parsers skip or scramble those.
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    disable_spellcheck(doc)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY
    pf = normal.paragraph_format
    pf.space_after = Pt(2)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    add_para(doc, NAME.upper(), size=20, bold=True, color=NAVY, align="center", after=1, space=1.0)
    add_para(
        doc,
        f"{LOCATION}  |  {PHONE}  |  {EMAIL}",
        size=10.5,
        align="center",
        after=1,
        space=1.0,
    )
    tag = add_para(doc, HEADLINE, size=10.5, italic=True, color=NAVY, align="center", after=2, space=1.0)
    add_bottom_border(tag, size="16")

    add_heading(doc, "Summary")
    add_para(doc, SUMMARY, size=10, after=2)

    add_heading(doc, "Skills")
    add_para(doc, SKILLS, size=10, after=2)

    add_heading(doc, "Experience")
    for job in JOBS:
        add_para(doc, job["title"], size=11, bold=True, color=NAVY, before=5, after=0, space=1.05)
        add_para(
            doc,
            f"{job['company']}, {job['location']}",
            size=10.5,
            bold=True,
            after=0,
            space=1.05,
        )
        add_para(doc, job["dates"], size=10, italic=True, color=MUTED, after=1, space=1.05)
        for bullet in job["bullets"]:
            add_bullet(doc, bullet)

    add_heading(doc, "Education")
    for item in EDUCATION:
        add_para(doc, item["school"], size=10.5, bold=True, color=NAVY, before=3, after=0, space=1.05)
        add_para(doc, item["detail"], size=10, after=1, space=1.05)

    add_heading(doc, "Patents")
    for patent in PATENTS:
        add_bullet(doc, patent)

    out = OUT_DIR / "Anthony_Ty_Marler_Indeed_Resume.docx"
    doc.save(out)
    return out


def build_markdown() -> Path:
    lines = [
        f"# {NAME}",
        "",
        f"{LOCATION} | {PHONE} | {EMAIL}",
        "",
        f"**{HEADLINE}**",
        "",
        "## Summary",
        "",
        SUMMARY,
        "",
        "## Skills",
        "",
        SKILLS,
        "",
        "## Experience",
        "",
    ]
    for job in JOBS:
        lines += [
            f"### {job['title']}",
            f"{job['company']}, {job['location']}",
            job["dates"],
            "",
        ]
        for bullet in job["bullets"]:
            lines.append(f"- {bullet}")
        lines.append("")
    lines += ["## Education", ""]
    for item in EDUCATION:
        lines += [f"**{item['school']}**", item["detail"], ""]
    lines += ["## Patents", ""]
    for patent in PATENTS:
        lines.append(f"- {patent}")
    lines.append("")
    out = OUT_DIR / "Anthony_Ty_Marler_Indeed_Resume.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def build_html() -> Path:
    job_html = []
    for job in JOBS:
        bullets = "\n".join(f"<li>{escape(b)}</li>" for b in job["bullets"])
        job_html.append(
            f"""<div class="job">
  <p class="title">{escape(job['title'])}</p>
  <p class="company">{escape(job['company'])}, {escape(job['location'])}</p>
  <p class="dates">{escape(job['dates'])}</p>
  <ul>
    {bullets}
  </ul>
</div>"""
        )
    edu_html = "\n".join(
        f"<p class=\"edu\"><b>{escape(item['school'])}</b><br />{escape(item['detail'])}</p>"
        for item in EDUCATION
    )
    patents = "\n".join(f"<li>{escape(p)}</li>" for p in PATENTS)
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <title>{escape(NAME)} - Indeed Resume</title>
  <style>
    @page {{ size: Letter; margin: 0.42in 0.58in 0.38in 0.58in; }}
    * {{ box-sizing: border-box; }}
    html, body {{
      margin: 0;
      padding: 0;
      background: #fff;
      color: #222;
      font-family: "Liberation Sans", Arial, Helvetica, sans-serif;
      font-size: 9.6pt;
      line-height: 1.16;
    }}
    .page {{ max-width: 8.5in; margin: 0 auto; }}
    h1 {{
      margin: 0;
      text-align: center;
      font-size: 20pt;
      letter-spacing: 0.03em;
      color: #1f3a5f;
    }}
    .contact, .headline {{
      text-align: center;
      margin: 2pt 0 0;
    }}
    .headline {{
      color: #1f3a5f;
      font-style: italic;
      padding-bottom: 5pt;
      border-bottom: 1.4pt solid #1f3a5f;
      margin-bottom: 2pt;
    }}
    h2 {{
      font-size: 11pt;
      letter-spacing: 0.05em;
      text-transform: uppercase;
      color: #1f3a5f;
      border-bottom: 0.9pt solid #1f3a5f;
      margin: 7pt 0 2pt;
      padding-bottom: 1pt;
    }}
    p {{ margin: 0 0 3pt; }}
    ul {{ margin: 1pt 0 3pt 1.05em; padding: 0; }}
    li {{ margin: 0 0 1pt; }}
    .job {{ margin-top: 5pt; }}
    .title {{ font-weight: 700; color: #1f3a5f; font-size: 11pt; margin: 0; }}
    .company {{ font-weight: 700; margin: 0; }}
    .dates {{ color: #4a4a4a; font-style: italic; margin: 0 0 2pt; }}
    .edu {{ margin: 3pt 0; }}
  </style>
</head>
<body>
  <div class="page">
    <h1>{escape(NAME.upper())}</h1>
    <p class="contact">{escape(LOCATION)} &nbsp;|&nbsp; {escape(PHONE)} &nbsp;|&nbsp; {escape(EMAIL)}</p>
    <p class="headline">{escape(HEADLINE)}</p>
    <h2>Summary</h2>
    <p>{escape(SUMMARY)}</p>
    <h2>Skills</h2>
    <p>{escape(SKILLS)}</p>
    <h2>Experience</h2>
    {"".join(job_html)}
    <h2>Education</h2>
    {edu_html}
    <h2>Patents</h2>
    <ul>
      {patents}
    </ul>
  </div>
</body>
</html>
"""
    out = OUT_DIR / "Anthony_Ty_Marler_Indeed_Resume.html"
    out.write_text(html, encoding="utf-8")
    return out


def assert_ats_safe(docx_path: Path) -> None:
    """Fail if the DOCX uses layouts Indeed / ATS commonly misread."""
    doc = Document(str(docx_path))
    problems = []
    if doc.tables:
        problems.append(f"contains {len(doc.tables)} table(s)")
    for section in doc.sections:
        if section.header.paragraphs and any(p.text.strip() for p in section.header.paragraphs):
            problems.append("header has text")
        if section.footer.paragraphs and any(p.text.strip() for p in section.footer.paragraphs):
            problems.append("footer has text")
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)
    joined_l = joined.lower()
    for needle in ("\t", "·", "–", "—"):
        if needle in joined:
            problems.append(f"contains special character {needle!r}")
    required = [NAME, PHONE, EMAIL, "Summary", "Skills", "Experience", "Education", "Patents"]
    for item in required:
        if item.lower() not in joined_l:
            problems.append(f"missing required text: {item}")
    # Linear order: Independent before Vistra before Army
    idx_ind = joined.find("Independent")
    idx_vis = joined.find("Vistra Corporate Services Company")
    idx_army = joined.find("U.S. Army")
    if not (0 <= idx_ind < idx_vis < idx_army):
        problems.append("experience is not reverse-chronological")
    if "February 1984 - February 1987" not in joined:
        problems.append("Army dates missing or not latest")
    if "Fractional CTO" not in joined:
        problems.append("headline missing Fractional CTO")
    if "internet television streaming application" not in joined:
        problems.append("missing latest Independent AI applications")
    if problems:
        raise SystemExit("ATS checks failed:\n- " + "\n- ".join(problems))
    print("ATS checks passed")


def main() -> None:
    md = build_markdown()
    html = build_html()
    docx_path = build_docx()
    assert_ats_safe(docx_path)
    print(f"Wrote {md}")
    print(f"Wrote {html}")
    print(f"Wrote {docx_path} ({docx_path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
