#!/usr/bin/env python3
"""Generate the NextEra Energy cover letter as a one-page DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BODY = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
OUT = Path(__file__).resolve().parent / "NextEra_Energy_Cover_Letter.docx"

PARAS = [
    "I am writing to apply to NextEra Energy. I have spent most of my career building the systems that sit between generation operations and the people who have to act on that data: traders, compliance officers, and plant leadership. At Vistra I designed the environmental and generation-data platform that accounted for $60 million in profit in its first two months, and I want to bring that same kind of work to NextEra's generation fleet and environmental markets.",
    "The result was not a trading strategy I invented on a desk. It was a visibility problem we finally solved. Actual environmental data, planned environmental data, and planned generation data lived in different systems, in different shapes, and on different clocks. Traders could not see a normalized picture of what the fleet had done, what it was planned to do, and what that implied for environmental position. I designed and built an AWS platform that ingested SCADA, PLC, and third-party environmental feeds into a single normalized store of actual and planned environmental data alongside planned generation. Once that picture was in front of the trading desk, they could see and capture value that had not been thinkable when the data was fragmented. Trading latency dropped from weeks to hours. Compliance reporting time fell by about 80 percent. In the first two months the platform accounted for $60 million in profit.",
    "That project sat inside a broader generation architecture practice. At Vistra I also delivered a data-driven process engine, mine-planning software that measured productivity by shift, and web applications and Tableau dashboards for mining, fossil generation, battery, and solar operations. Before that I led Fossil Applications for Luminant generation and mining at Accenture, and at Capgemini I architected the MSHA, emissions, GIS, mobility, PI, and SharePoint systems those plants ran on. I am a named inventor on two U.S. patents and have spent the last year shipping AI-assisted products with Cursor, Claude, and ChatGPT.",
    "NextEra is the place that combination belongs. You generate at a scale where environmental position, planned output, and market action have to move together. I know how to make that data trustworthy, timely, and usable by the people who can turn it into results. I would welcome the chance to discuss how I can do that on your team.",
]


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=BODY):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def add_para(doc, text, *, size=11, bold=False, italic=False, color=BODY, align="left",
             before=0, after=8, space=1.12):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = space
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    add_para(doc, "ANTHONY TY MARLER", size=16, bold=True, color=NAVY, align="center", after=2, space=1.0)
    add_para(
        doc,
        "Winnsboro, TX 75494  |  (903) 474-3647  |  tymarler@gmail.com",
        size=11,
        align="center",
        after=10,
        space=1.0,
    )
    add_para(doc, "September 23, 2026", size=11, after=10)
    add_para(doc, "Hiring Team", size=11, after=0, space=1.08)
    add_para(doc, "NextEra Energy", size=11, after=0, space=1.08)
    add_para(doc, "Juno Beach, Florida", size=11, after=10, space=1.08)
    add_para(doc, "Dear Hiring Team:", size=11, after=8)

    for text in PARAS:
        add_para(doc, text, size=11, after=8)

    add_para(doc, "Sincerely,", size=11, after=14)
    add_para(doc, "Anthony Ty Marler", size=11, bold=True, color=NAVY, after=0, space=1.08)
    add_para(doc, "(903) 474-3647  |  tymarler@gmail.com", size=11, color=MUTED, after=0, space=1.08)

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
